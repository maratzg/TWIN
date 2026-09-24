import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.input_engine.mapping.service import MappingService, normalize_header


@dataclass
class ParsedTable:
    source_type: str
    sheet_names: list[str]
    selected_sheet: str | None
    header_row: int
    columns: list[str]
    rows: list[dict[str, Any]]
    samples: dict[str, list[str]]


def json_safe(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _unique_headers(values: list[Any]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value not in (None, "") else f"COLUMN_{index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return headers


class TabularParser:
    def __init__(self) -> None:
        self.known_aliases = set(MappingService().alias_to_field)

    def parse(self, filename: str, content: bytes) -> ParsedTable:
        suffix = Path(filename).suffix.lower()
        if suffix == ".csv":
            return self._parse_csv(content)
        if suffix == ".xlsx":
            return self._parse_excel(content)
        raise ValueError("Unsupported file type. Upload a .csv or .xlsx file.")

    def _parse_csv(self, content: bytes) -> ParsedTable:
        text = self._decode_csv(content)
        delimiter = self._detect_delimiter(text)
        matrix = [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
        header_index = self._detect_header_index(matrix)
        columns, rows = self._rows_from_matrix(matrix, header_index)
        return ParsedTable("csv", [], None, header_index + 1, columns, rows, self._samples(columns, rows))

    def _parse_excel(self, content: bytes) -> ParsedTable:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheets: list[tuple[str, list[list[Any]], int]] = []
        for worksheet in workbook.worksheets:
            matrix = [list(row) for row in worksheet.iter_rows(values_only=True)]
            if not matrix:
                continue
            header_index = self._detect_header_index(matrix)
            populated = sum(any(value not in (None, "") for value in row) for row in matrix[header_index + 1 :])
            sheets.append((worksheet.title, matrix, populated))
        if not sheets:
            raise ValueError("Workbook contains no readable data.")
        selected_name, matrix, _ = max(sheets, key=lambda item: item[2])
        header_index = self._detect_header_index(matrix)
        columns, rows = self._rows_from_matrix(matrix, header_index)
        return ParsedTable(
            "xlsx",
            [worksheet.title for worksheet in workbook.worksheets],
            selected_name,
            header_index + 1,
            columns,
            rows,
            self._samples(columns, rows),
        )

    @staticmethod
    def _decode_csv(content: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp1252"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError("CSV encoding is not supported.")

    @staticmethod
    def _detect_delimiter(text: str) -> str:
        lines = [line for line in text.splitlines()[:20] if line.strip()]
        scored = []
        for delimiter in (",", ";", "\t", "|"):
            counts = [line.count(delimiter) for line in lines]
            populated = [count for count in counts if count > 0]
            consistency = max((populated.count(count) for count in set(populated)), default=0)
            scored.append((consistency, sum(populated), delimiter))
        consistency, total, delimiter = max(scored)
        return delimiter if consistency and total else ","

    def _detect_header_index(self, matrix: list[list[Any]]) -> int:
        if not matrix:
            raise ValueError("File contains no rows.")
        candidates: list[tuple[float, int]] = []
        for index, row in enumerate(matrix[:20]):
            values = [value for value in row if value not in (None, "")]
            if len(values) < 2:
                candidates.append((-1.0, index))
                continue
            string_count = sum(isinstance(value, str) for value in values)
            unique_count = len({str(value).strip().upper() for value in values})
            alias_hits = sum(normalize_header(str(value)) in self.known_aliases for value in values)
            score = string_count + unique_count * 0.25 + alias_hits * 4 - index * 0.01
            candidates.append((score, index))
        return max(candidates)[1]

    @staticmethod
    def _rows_from_matrix(
        matrix: list[list[Any]], header_index: int
    ) -> tuple[list[str], list[dict[str, Any]]]:
        columns = _unique_headers(matrix[header_index])
        rows: list[dict[str, Any]] = []
        for raw_row in matrix[header_index + 1 :]:
            padded = list(raw_row) + [None] * max(0, len(columns) - len(raw_row))
            if not any(value not in (None, "") for value in padded[: len(columns)]):
                continue
            rows.append({column: json_safe(padded[index]) for index, column in enumerate(columns)})
        return columns, rows

    @staticmethod
    def _samples(columns: list[str], rows: list[dict[str, Any]]) -> dict[str, list[str]]:
        result: dict[str, list[str]] = {}
        for column in columns:
            seen: list[str] = []
            for row in rows:
                value = row.get(column)
                if value in (None, ""):
                    continue
                text = str(value)
                if text not in seen:
                    seen.append(text)
                if len(seen) == 3:
                    break
            result[column] = seen
        return result
