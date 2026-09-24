from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "demo-data" / "company_a_shipments.xlsx"


def main() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Shipments Sep"
    sheet.append(["Company A operational shipment export — September 2026"])
    headers = ["CUST", "REF", "B/L", "CONT", "VSL", "VOY", "POL", "POD", "ETD", "ETA", "GWT", "TEU", "REV", "CURRENCY", "STATUS"]
    sheet.append(headers)
    rows = [
        ["IKEA", "RTM-8821", "BL-A-1001", "MSCU1234567", "Ever Given", "EG42", "NLRTM", "LVLPX", "2026-09-01", "2026-09-10", 24500, 2, 12450.00, "EUR", "Booked"],
        ["ABB", "HAM-4420", "BL-A-1002", "CMAU7654321", "CMA CGM Titan", "CG11", "DEHAM", "SEGOT", "03/09/2026", "12/09/2026", 18750, 1, 9800.50, "EUR", "In transit"],
        ["Bosch", "ANR-1184", "BL-A-1003", "TLLU4455667", "Maersk Lima", "ML8", "BEANR", "FIKTK", "2026-09-04", "2026-09-14", "18,4t", 1, 7500, "EUR", "Booked"],
        ["IKEA", "RTM-8821", "BL-A-1001", "MSCU1234567", "Ever Given", "EG42", "NLRTM", "LVLPX", "2026-09-01", "2026-09-10", 24500, 2, 12450.00, "EUR", "Booked"],
        ["Nordic Parts", "GDN-7005", "BL-A-1005", "BAD-CONTAINER", "OOCL Europe", "OE5", "PLGDN", "LVLPX", "2026-09-20", "2026-09-18", 15400, 1, 6100, "EUR", "Booked"],
        [None, None, None, None, "MSC Aurora", "A12", "DKCPH", "LVLPX", "not-a-date", None, None, None, None, "EUR", "Draft"],
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D3D69")
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:O{sheet.max_row}"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = max(12, min(28, max(len(str(cell.value or "")) for cell in column) + 2))

    notes = workbook.create_sheet("Read me")
    notes.append(["Synthetic demo data only"])
    notes.append(["Includes REF mapping review, a value correction (18,4t), a duplicate, missing identifiers, and a title row for header detection."])
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
